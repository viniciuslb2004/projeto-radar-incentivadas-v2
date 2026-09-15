"""Gera o .xlsx formatado exportado pela aba Busca (ver webapp/main.py::busca_exportar).

As linhas exportadas vem PRONTAS do front-end (ultimosResultados, ja renderizado na tela) --
o backend nunca re-roda a busca aqui, so monta a planilha em cima do que foi passado. Isso
garante que o arquivo bate exatamente com o que a pessoa viu, mesmo que o banco mude entre o
clique em "Buscar" e o clique em "Exportar".
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# Mesmas colunas/rotulos do CSV que este export substitui (ver busca.js, funcao renderResultados).
COLUNAS = [
    ("cliente", "Cliente"),
    ("cnpj", "CNPJ"),
    ("agencia", "Agência"),
    ("setor_bndes", "Setor"),
    ("subsetor_bndes", "Subsetor"),
    ("segmento", "Segmento"),
    ("uf", "UF"),
    ("data_contratacao", "Data"),
    ("valor_contratado", "Valor contratado"),
    ("valor_desembolsado", "Valor desembolsado"),
    ("descricao_projeto", "Descrição do projeto"),
    ("score", "Similaridade"),
    ("motivo", "Motivo da correspondência"),
]

_COLUNAS_MOEDA = {"valor_contratado", "valor_desembolsado"}
_LARGURA_MAXIMA = 60

# Cores de marca (ver webapp/static/css/style.css :root) -- reaproveitadas aqui em vez de
# inventar uma paleta nova pro arquivo exportado.
_NAVY = "223850"
_BORDA = "EBECED"

# Colunas do export de Transacoes Salvas (ver webapp/salvos.py::listar_operacoes_salvas) --
# mesma base de COLUNAS acima, so troca "score"/"motivo" (conceitos da Busca, sem
# sentido aqui) por "nota"/"salvo_em" (proprios de uma operacao favoritada).
COLUNAS_SALVAS = [c for c in COLUNAS if c[0] not in ("score", "motivo")] + [
    ("nota", "Nota pessoal"),
    ("salvo_em", "Salvo em"),
]


def _gerar_xlsx(titulo_aba: str, colunas: list, linhas: list) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = titulo_aba

    fonte_header = Font(color="FFFFFF", bold=True)
    preenchimento_header = PatternFill(start_color=_NAVY, end_color=_NAVY, fill_type="solid")
    borda_fina = Border(*(Side(style="thin", color=_BORDA) for _ in range(4)))

    ws.append([rotulo for _, rotulo in colunas])
    for cel in ws[1]:
        cel.font = fonte_header
        cel.fill = preenchimento_header
        cel.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"

    larguras = [len(rotulo) for _, rotulo in colunas]
    for linha in linhas:
        valores = [linha.get(chave) for chave, _ in colunas]
        ws.append(valores)
        r = ws.max_row
        for i, (chave, _) in enumerate(colunas, start=1):
            cel = ws.cell(row=r, column=i)
            cel.border = borda_fina
            if chave in _COLUNAS_MOEDA and cel.value is not None:
                cel.number_format = "R$ #,##0.00"
            texto = str(cel.value) if cel.value is not None else ""
            larguras[i - 1] = min(_LARGURA_MAXIMA, max(larguras[i - 1], len(texto)))

    for i, largura in enumerate(larguras, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = largura + 2

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def gerar_xlsx_busca(query: str, linhas: list) -> bytes:
    return _gerar_xlsx("Busca", COLUNAS, linhas)


def gerar_xlsx_operacoes_salvas(linhas: list) -> bytes:
    """Exporta as operacoes salvas do usuario (ver webapp/main.py::salvos_exportar) --
    mesmo gerador/estilo do export da Busca, colunas adaptadas (ver COLUNAS_SALVAS)."""
    return _gerar_xlsx("Transações Salvas", COLUNAS_SALVAS, linhas)
